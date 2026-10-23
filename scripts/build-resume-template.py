from pathlib import Path
from docx import Document
from docx.shared import Pt, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

doc=Document(); sec=doc.sections[0]
sec.page_width=Mm(210); sec.page_height=Mm(297)
sec.top_margin=Pt(30); sec.bottom_margin=Pt(30); sec.left_margin=Pt(36); sec.right_margin=Pt(36)
style=doc.styles['Normal']; style.font.name='Times New Roman'; style.font.size=Pt(11)
style.paragraph_format.space_after=Pt(2); style.paragraph_format.line_spacing=1
for sn in ['Title','Heading 1','List Bullet']:
 st=doc.styles[sn]; st.font.name='Times New Roman'; st.font.color.rgb=__import__('docx').shared.RGBColor(0,0,0)

def para(text='',bold=False,italic=False):
 p=doc.add_paragraph();r=p.add_run(text);r.bold=bold;r.italic=italic;return p

def row(left,right,bold=False,italic=False):
 p=para();p.paragraph_format.tab_stops.add_tab_stop(Pt(523),WD_TAB_ALIGNMENT.RIGHT)
 r=p.add_run(left);r.bold=bold;r.italic=italic
 r=p.add_run('\t'+right);r.italic=italic
 p.paragraph_format.keep_with_next=True
 return p

def heading(text):
 p=para(text,True);p.paragraph_format.space_before=Pt(7);p.paragraph_format.space_after=Pt(3);p.paragraph_format.keep_with_next=True
 borders=OxmlElement('w:pBdr');line=OxmlElement('w:bottom')
 for k,v in [('val','single'),('sz','4'),('space','2'),('color','000000')]:line.set(qn('w:'+k),v)
 borders.append(line);p._p.get_or_add_pPr().append(borders)

def bullet(lead,text):
 p=doc.add_paragraph(style='List Bullet');p.paragraph_format.left_indent=Pt(10);p.paragraph_format.first_line_indent=Pt(-10)
 p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1
 p.add_run(lead+' ').bold=True;p.add_run(text)

p=doc.add_paragraph(style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(4)
r=p.add_run('[YOUR NAME]');r.bold=True;r.font.size=Pt(20)
p=para('[City, Country] | [Phone] | [Email]');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p=para('[LinkedIn or portfolio URL]');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
heading('PROFILE')
para('[One sentence connecting your strongest experience and skills to the target role.]')
p=para();p.add_run('Skills: ').bold=True;p.add_run('[Relevant tools, methods and skills supported by your experience]')
p=para();p.add_run('Languages: ').bold=True;p.add_run('[Languages and proficiency levels]')
heading('EDUCATION')
row('[University / Graduate School]','[City, Country]',True)
row('[Master\'s degree and field]','[Start - End / Expected]',italic=True)
para('[Relevant research, modules or honours - include only what applies]')
row('[University / Business School]','[City, Country]',True)
row('[Bachelor\'s degree and field]','[Start - End]',italic=True)
para('[Relevant specialisation, honours or GPA - include only what applies]')
heading('PROFESSIONAL EXPERIENCE')
for i in range(1,4):
 p=row(f'[Employer {i} / Team]','[City, Country]',True)
 if i>1:p.paragraph_format.space_before=Pt(5)
 row('[Official job title]','[Month Year - Month Year]',italic=True)
 bullet('[Action / contribution]:','[Describe the work, your contribution and a supported outcome.]')
 bullet('[Method / delivery]:','[Explain how you used relevant tools or skills to deliver a result.]')
 bullet('[Impact / improvement]:','[State a verified outcome or scale; add a metric only if supported.]')
heading('EXTRACURRICULAR ACTIVITIES')
for i in range(1,3):
 p=row(f'[Activity {i} / Organisation]','[City, Country]',True)
 if i>1:p.paragraph_format.space_before=Pt(5)
 row('[Role / Position]','[Month Year - Month Year]',italic=True)
 bullet('[Initiative / leadership]:','[Describe your contribution and what you delivered.]')
 bullet('[Outcome / collaboration]:','[Describe a supported result, reach or team contribution.]')
doc.core_properties.author='Tailor';doc.core_properties.title='One Page Resume Template'
doc.core_properties.subject='Blank resume template with two education entries, three roles and two activities'
# Remove built-in title theme styling to match the serif reference exactly.
for st in doc.styles:
 for rf in st.element.xpath('.//w:rFonts'):
  for key in list(rf.attrib):
   if 'theme' in key.lower(): del rf.attrib[key]
for border in doc.styles['Title'].element.xpath('.//w:pBdr'):
 border.getparent().remove(border)
for p in doc.paragraphs:
 for run in p.runs: run.font.name='Times New Roman'
doc.save(str(Path(__file__).resolve().parents[1] / 'public/templates/classic-resume.docx'))
